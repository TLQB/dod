from django.shortcuts import render
from django.http import JsonResponse, HttpResponse
from .forms import DataProcessingForm, DataUpdateDiscountForm
import pandas as pd
from .utils import (
    get_sheet_names,
    get_all_brands,
    get_column_names,
    process_and_export_data,
)
import tempfile
import os
import io
from openpyxl import load_workbook


def process_excel_data(request):
    if request.method == "POST":
        form = DataProcessingForm(request.POST, request.FILES)

        input_file = request.FILES["input_filename"]
        form.fields["sheet_name"].choices = [
            (sheet, sheet) for sheet in get_sheet_names(input_file)
        ]
        form.fields["brand"].choices = [
            (brand, brand) for brand in get_all_brands(input_file)
        ]
        form.fields["column"].choices = [
            (col, col) for col in get_column_names(input_file)
        ]

        if form.is_valid():
            print(form.cleaned_data)
            input_file = request.FILES["input_filename"]
            output_filename = form.cleaned_data["output_filename"]
            sheet_names = form.cleaned_data["sheet_name"]
            brands = form.cleaned_data["brand"]
            columns = form.cleaned_data["column"]

            # Tạo tệp tạm thời để lưu kết quả
            with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_file:
                output_file_path = tmp_file.name

            # try:
            process_and_export_data(
                input_file, output_file_path, columns, sheet_names, brands
            )

            # Đọc tệp kết quả và trả về cho người dùng
            with open(output_file_path, "rb") as file:
                response = HttpResponse(
                    file.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                response["Content-Disposition"] = (
                    f'attachment; filename="{output_filename}"'
                )

            # Xóa tệp tạm sau khi đã gửi
            os.unlink(output_file_path)

            return response
            # except Exception as e:
            #     return JsonResponse({'error': str(e)}, status=400)
    else:
        form = DataProcessingForm()

    return render(request, "tools_template/export_data_excel.html", {"form": form})


def get_file_data(request):
    if request.method == "POST" and request.FILES:
        form = DataProcessingForm()
        input_file = request.FILES["input_filename"]
        sheets_choices = [(sheet, sheet) for sheet in get_sheet_names(input_file)]
        brands_choices = [(brand, brand) for brand in get_all_brands(input_file)]
        columns_choices = [(col, col) for col in get_column_names(input_file)]

        data = {
            "sheets_choices": sheets_choices,
            "brands_choices": brands_choices,
            "columns_choices": columns_choices,
        }

        return JsonResponse(data)
    return JsonResponse({"error": "Invalid request"}, status=400)


def process_files(file_data, file_update):
    # Đọc dữ liệu từ input file
    df_input = pd.read_excel(file_data, sheet_name="MD4 Special DC", header=2)
    # Tạo dictionary từ df_input với Product Code là khóa
    # và Discount Rate là giá trị
    if "Product Code" in df_input.columns and "Discount Rate" in df_input.columns:
        df_input["Product Code"] = (
            pd.to_numeric(df_input["Product Code"], errors="coerce")
            .fillna(0)
            .astype(int)
        )
        df_input["Discount Rate"] = pd.to_numeric(
            df_input["Discount Rate"], errors="coerce"
        ).fillna(0)
        res = dict(zip(df_input["Product Code"], df_input["Discount Rate"]))

    # Đọc file Excel đích với multi-index header
    df = pd.read_excel(file_update, sheet_name="sheet1", header=[0, 1])

    # Hàm tìm key cho Product Code trong DataFrame
    def get_key_product_code(df_output):
        for key in df_output.keys():
            if "Product Code" in key:
                return key
        return None

    # Hàm tìm key cho Offline Discount Rate trong DataFrame
    def get_key_rate_dc(df_output):
        for key in df_output.keys():
            if "Offline Discount Rate" in key and "D/C Rate" in key:
                return key
        return None

    # Lấy các khóa để truy xuất cột
    product_code_key = get_key_product_code(df)
    dc_rate_key = get_key_rate_dc(df)

    # Mở workbook gốc với openpyxl
    wb = load_workbook(file_update)
    ws = wb["sheet1"]

    # Ghi từng giá trị cập nhật vào file Excel theo index của DataFrame
    for idx, row in df.iterrows():
        product_code_value = row[product_code_key]
        if product_code_value in res and res[product_code_value] != 0:
            # Tìm vị trí (row, col) trong Excel để ghi
            col_index = df.columns.get_loc(dc_rate_key) + 1
            # Cột trong Excel bắt đầu từ 1
            excel_row_index = (
                idx + 3
            )  # Cộng thêm 3 do header chiếm 2 dòng và data bắt đầu từ dòng 3
            ws.cell(
                row=excel_row_index, column=col_index, value=res[product_code_value]
            )

    # Tạo một buffer trong bộ nhớ
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return output


def update_discount_rate(request):
    if request.method == "POST":
        form = DataUpdateDiscountForm(request.POST, request.FILES)

        if form.is_valid():
            file_data = request.FILES["input_filename"]
            file_update = request.FILES["output_filename"]

            # Gọi hàm xử lý file
            output = process_files(file_data, file_update)

            # Tạo HTTP response với file Excel
            response = HttpResponse(
                output.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = (
                "attachment; filename=updated_output_file.xlsx"
            )

            return response

    else:
        form = DataUpdateDiscountForm()

    return render(request, "tools_template/update_discount_rate.html", {"form": form})
